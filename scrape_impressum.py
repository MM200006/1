#!/usr/bin/env python3
"""
Liest Impressum-Seiten aus und extrahiert Kontaktdaten
(Telefon, E-Mail, Geschäftsführer) automatisch.
Aktualisiert die Excel-Kontaktliste.

Verwendung:
    python scrape_impressum.py                    # Alle Einträge ohne Telefon
    python scrape_impressum.py --row 5            # Nur Zeile 5
    python scrape_impressum.py --max 10           # Maximal 10 Einträge
"""

import argparse
import re
import time
import sys
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

EXCEL_PATH = "kontaktliste_akquise.xlsx"

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                  "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
    "Accept-Language": "de-DE,de;q=0.9,en;q=0.8",
}


def extract_phone(text):
    patterns = [
        r"(?:Tel(?:efon)?|Phone|Fon|Ruf)[\s.:]*([+\d][\d\s/\-().]{7,20})",
        r"(\+49[\s\d/\-().]{8,20})",
        r"(0\d{2,4}[\s/\-][\d\s/\-]{5,15})",
    ]
    for pattern in patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            phone = match.group(1).strip()
            phone = re.sub(r"\s+", " ", phone)
            return phone
    return ""


def extract_email(text):
    match = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
    return match.group(0) if match else ""


def extract_ceo(text):
    patterns = [
        r"(?:Geschäftsführ(?:er|ung)|Vorstand|CEO|Managing Director|Inhaber|Vertretungsberechtig)[\s.:]*\n?\s*([A-ZÄÖÜ][a-zäöüß]+\s[A-ZÄÖÜ][a-zäöüß\-]+(?:\s[A-ZÄÖÜ][a-zäöüß\-]+)?)",
        r"(?:Geschäftsführ(?:er|ung)|CEO|Inhaber)[\s.:]*\n?\s*((?:(?:Dr\.|Prof\.|Dipl\.[\w.-]*)\s+)?[A-ZÄÖÜ][a-zäöüß]+\s[A-ZÄÖÜ][a-zäöüß\-]+)",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return match.group(1).strip()
    return ""


def scrape_impressum(url):
    result = {"phone": "", "email": "", "ceo": ""}
    try:
        resp = requests.get(url, headers=HEADERS, timeout=10, allow_redirects=True)
        resp.raise_for_status()
        soup = BeautifulSoup(resp.text, "html.parser")

        for tag in soup(["script", "style", "nav", "header", "footer"]):
            tag.decompose()

        text = soup.get_text(separator="\n")
        text = re.sub(r"\n{3,}", "\n\n", text)

        result["phone"] = extract_phone(text)
        result["email"] = extract_email(text)
        result["ceo"] = extract_ceo(text)

    except requests.RequestException as e:
        print(f"  Fehler bei {url}: {e}")

    return result


def try_impressum_urls(base_url):
    base = base_url.rstrip("/")
    candidates = [
        base + "/impressum",
        base + "/impressum/",
        base + "/de/impressum",
        base + "/imprint",
        base + "/legal",
        base + "/legal-notice",
    ]
    for url in candidates:
        try:
            resp = requests.head(url, headers=HEADERS, timeout=5, allow_redirects=True)
            if resp.status_code == 200:
                return url
        except requests.RequestException:
            continue
    return candidates[0]


def update_excel(target_row=None, max_entries=None):
    wb = load_workbook(EXCEL_PATH)
    ws = wb["Kontaktliste Akquise"]

    count = 0
    for row in range(2, ws.max_row + 1):
        if target_row and row != target_row:
            continue

        phone_cell = ws.cell(row=row, column=6).value
        if phone_cell and not target_row:
            continue

        website = ws.cell(row=row, column=4).value
        name = ws.cell(row=row, column=3).value
        if not website:
            continue

        if max_entries and count >= max_entries:
            break

        impressum_url = ws.cell(row=row, column=5).value
        if not impressum_url or impressum_url == website.rstrip("/") + "/impressum":
            impressum_url = try_impressum_urls(website)
            ws.cell(row=row, column=5, value=impressum_url)

        print(f"[{row-1}] {name}: {impressum_url}")
        data = scrape_impressum(impressum_url)

        if data["phone"]:
            ws.cell(row=row, column=6, value=data["phone"])
            print(f"  Tel: {data['phone']}")
        if data["email"]:
            ws.cell(row=row, column=7, value=data["email"])
            print(f"  E-Mail: {data['email']}")
        if data["ceo"]:
            ws.cell(row=row, column=8, value=data["ceo"])
            print(f"  GF: {data['ceo']}")

        if not any(data.values()):
            print("  Keine Daten gefunden")

        count += 1
        time.sleep(1.5)

    wb.save(EXCEL_PATH)
    print(f"\nFertig! {count} Einträge verarbeitet. Datei gespeichert: {EXCEL_PATH}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Impressum-Daten scrapen und Excel aktualisieren")
    parser.add_argument("--row", type=int, help="Nur bestimmte Zeile verarbeiten")
    parser.add_argument("--max", type=int, help="Maximale Anzahl zu verarbeitender Einträge")
    args = parser.parse_args()
    update_excel(target_row=args.row, max_entries=args.max)
